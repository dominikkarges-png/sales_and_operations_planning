"""
S&OP Prognose – Streamlit App (Lemken)
Long-Format Eingabe: Datum, Monat, Produktgruppe, Phase, Menge, ...
Sheet: Long_Monat_Produktgruppe (oder bestes Match per Auto-Detection)

Features:
- Filter (Produktgruppe, Jahr, Monat, Phase)
- KPIs + Donut (Produktgruppe, kompakt)
- Liniendiagramm Ist (navy) / Forecast (grau) mit Range-Slider
- Editierbare Pivot-Tabelle: 3 Monate Rueckblick + alle zukuenftigen Monate
- Freigabe-Workflow Planer/Reviewer mit Session-State-Persistenz (Cloud-kompatibel)
- Excel-Export der freigegebenen Prognose

Start:
    pip install streamlit pandas plotly openpyxl
    streamlit run streamlit_app.py
"""

import base64
from datetime import datetime
from io import BytesIO
from pathlib import Path
from zoneinfo import ZoneInfo

import pandas as pd
import plotly.graph_objects as go
import plotly.express as px
import streamlit as st
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ============================================================
# Konfiguration
# ============================================================
BASE_DIR    = Path(__file__).parent
LOGO_FILE   = BASE_DIR / "logo.png"
LEMKEN_FILE = BASE_DIR / "LemkenGmbH_Logo.png"

PRIMARY       = "#000055"
BLUE_MEDIUM   = "#6688BB"
BLUE_LIGHT    = "#BBCCEE"
BLUE_LIGHTEST = "#E4E4F6"
GRAY_MEDIUM   = "#999999"
GRAY_LIGHT    = "#D2D2D2"
FORECAST_GRAY = "#B0B0B0"

ACCENT_ORANGE     = "#FFC000"
ACCENT_DARK_GREEN = "#00B050"
ACCENT_RED        = "#C31924"

WP_PALETTE = [PRIMARY, BLUE_MEDIUM, BLUE_LIGHT, GRAY_MEDIUM, GRAY_LIGHT, BLUE_LIGHTEST]
LOGO_MAXH = 80

STATUS_COLOR = {
    "ENTWURF":      GRAY_MEDIUM,
    "ZUR_PRUEFUNG": ACCENT_ORANGE,
    "FREIGEGEBEN":  ACCENT_DARK_GREEN,
    "ABGELEHNT":    ACCENT_RED,
}
STATUS_LABEL = {
    "ENTWURF":      "Entwurf",
    "ZUR_PRUEFUNG": "Zur Prüfung",
    "FREIGEGEBEN":  "Freigegeben",
    "ABGELEHNT":    "Abgelehnt",
}

COL_DATUM = "Datum"
COL_MONAT = "Monat"
COL_PG    = "Produktgruppe"
COL_PHASE = "Phase"
COL_MENGE = "Menge"

PHASE_IST      = "Ist"
PHASE_FORECAST = "Forecast"

KEY_COLS    = [COL_PG]
DB_KEY_COLS = ["produktgruppe"]
DB_TO_COL   = dict(zip(DB_KEY_COLS, KEY_COLS))

COLUMN_ALIASES = {
    COL_DATUM: ["datum", "date", "period", "zeitpunkt"],
    COL_MONAT: ["monat", "month", "periode", "period_label"],
    COL_PG:    ["produktgruppe", "product group", "productgroup", "pg",
                "produkt gruppe", "produkt-gruppe", "produkt"],
    COL_PHASE: ["phase", "typ", "type", "kategorie", "category"],
    COL_MENGE: ["menge", "quantity", "qty", "absatz", "absatzmenge",
                "value", "wert"],
}

PREFERRED_SHEET = "Long_Monat_Produktgruppe"
LOOKBACK_MONTHS = 3

LOCAL_TZ = ZoneInfo("Europe/Berlin")


def now_local() -> datetime:
    return datetime.now(LOCAL_TZ).replace(tzinfo=None)

st.set_page_config(page_title="S&OP Prognose", page_icon="📈", layout="wide")

# ============================================================
# Passwort-Schutz – muss vor allem anderen laufen
# ============================================================
PASSWORD = "wieselhuberNLP2026!"

if not st.session_state.get("authenticated", False):
    st.markdown(
        f"""
        <style>
        .login-box {{
            max-width: 400px; margin: 100px auto; padding: 40px;
            border: 2px solid #000055; border-radius: 8px;
            background: #FFFFFF; text-align: center;
        }}
        .login-box h2 {{ color: #000055; margin-bottom: 24px; }}
        </style>
        <div class="login-box">
            <h2>🔒 S&OP Prognose</h2>
        </div>
        """,
        unsafe_allow_html=True,
    )
    pw = st.text_input("Passwort", type="password", label_visibility="collapsed",
                       placeholder="Passwort eingeben...")
    if st.button("Anmelden", use_container_width=False):
        if pw == PASSWORD:
            st.session_state["authenticated"] = True
            st.rerun()
        else:
            st.error("Falsches Passwort.")
    st.stop()


# ============================================================
# Hilfsfunktionen
# ============================================================
def fmt_int(n) -> str:
    try:
        return f"{int(round(float(n))):,}".replace(",", ".")
    except (ValueError, TypeError):
        return "0"


def _clean_colname(name) -> str:
    if name is None:
        return ""
    s = str(name)
    for ch in ("\ufeff", "\xa0", "\u00ad", "\u200b", "\u200c", "\u200d"):
        s = s.replace(ch, " ")
    return " ".join(s.split())


def normalize_phase(val) -> str:
    if pd.isna(val):
        return "Unbekannt"
    s = str(val).strip().lower()
    if s in ("ist", "actual", "actuals", "historie", "historisch", "ist-daten"):
        return PHASE_IST
    if s in ("forecast", "prognose", "fc", "plan", "forecast-daten"):
        return PHASE_FORECAST
    return str(val).strip()


def resolve_columns(df_raw: pd.DataFrame) -> pd.DataFrame:
    cleaned       = {col: _clean_colname(col) for col in df_raw.columns}
    cleaned_lower = {col: cleaned[col].lower() for col in df_raw.columns}
    rename_map = {}
    for target, aliases in COLUMN_ALIASES.items():
        match = None
        for orig, clean in cleaned.items():
            if clean == target:
                match = orig
                break
        if match is None:
            alias_set = {a.lower() for a in aliases}
            for orig, clean_low in cleaned_lower.items():
                if clean_low in alias_set:
                    match = orig
                    break
        if match is not None:
            rename_map[match] = target
    df = df_raw.rename(columns=rename_map).copy()
    required = [COL_DATUM, COL_PG, COL_PHASE, COL_MENGE]
    missing = [c for c in required if c not in df.columns]
    if missing:
        gefunden = "\n".join(f"  - '{c}'  (gereinigt: '{cleaned[c]}')" for c in df_raw.columns)
        raise ValueError(
            f"Fehlende Spalten: {missing}\n\n"
            f"In der Datei gefundene Spalten:\n{gefunden}\n\n"
            f"Akzeptierte Aliase:\n"
            + "\n".join(f"  - {m}: {COLUMN_ALIASES[m]}" for m in missing)
        )
    return df


def prepare_df(df_raw: pd.DataFrame) -> pd.DataFrame:
    df = resolve_columns(df_raw)
    df[COL_DATUM] = pd.to_datetime(df[COL_DATUM], errors="coerce")
    df = df.dropna(subset=[COL_DATUM])
    if df.empty:
        raise ValueError("Keine gueltigen Datumswerte in Spalte 'Datum' gefunden.")
    df[COL_PG]    = df[COL_PG].fillna("(unbekannt)").astype(str).str.strip()
    df[COL_PHASE] = df[COL_PHASE].apply(normalize_phase)
    df[COL_MENGE] = pd.to_numeric(df[COL_MENGE], errors="coerce").fillna(0)
    df["Jahr"]    = df[COL_DATUM].dt.year
    df["MonatNr"] = df[COL_DATUM].dt.month
    return df


# ============================================================
# Excel-Loader mit Sheet- & Header-Detection
# ============================================================
def _alias_set() -> set:
    s = set()
    for target, aliases in COLUMN_ALIASES.items():
        s.add(target.lower())
        for a in aliases:
            s.add(a.lower())
    return s


def _pick_sheet(source) -> str:
    xls    = pd.ExcelFile(source)
    sheets = xls.sheet_names
    if PREFERRED_SHEET in sheets:
        return PREFERRED_SHEET
    valid_names = _alias_set()
    best_sheet, best_hits = sheets[0], -1
    for sh in sheets:
        try:
            preview = pd.read_excel(xls, sheet_name=sh, header=None, nrows=10)
        except Exception:
            continue
        hits = 0
        for i in range(len(preview)):
            row_vals = [_clean_colname(v).lower() for v in preview.iloc[i].tolist()]
            hits = max(hits, sum(1 for v in row_vals if v in valid_names))
        if hits > best_hits:
            best_hits, best_sheet = hits, sh
    return best_sheet


def _detect_header_row(raw: pd.DataFrame, max_scan: int = 10) -> int:
    valid_names = _alias_set()
    best_row, best_hits = 0, 0
    for i in range(min(max_scan, len(raw))):
        row_vals = [_clean_colname(v).lower() for v in raw.iloc[i].tolist()]
        hits = sum(1 for v in row_vals if v in valid_names)
        if hits > best_hits:
            best_hits, best_row = hits, i
    return best_row


def read_excel_smart(source) -> pd.DataFrame:
    sheet = _pick_sheet(source)
    if hasattr(source, "seek"):
        source.seek(0)
    raw = pd.read_excel(source, sheet_name=sheet, header=None)
    header_row = _detect_header_row(raw)
    if hasattr(source, "seek"):
        source.seek(0)
    return pd.read_excel(source, sheet_name=sheet, header=header_row)


# ============================================================
# Session State – DB-Ersatz (statt SQLite)
# ============================================================
def _init_session_db():
    if "db_forecast" not in st.session_state:
        st.session_state["db_forecast"] = pd.DataFrame(columns=[
            "produktgruppe", "monat", "menge", "status",
            "ersteller", "kommentar", "created_at", "updated_at"
        ])
    if "db_history" not in st.session_state:
        st.session_state["db_history"] = pd.DataFrame(columns=[
            "id", "produktgruppe", "monat", "menge_alt", "menge_neu",
            "differenz", "status", "user", "kommentar", "aktion", "timestamp"
        ])

_init_session_db()


def save_forecast_bulk(df_long: pd.DataFrame, original_lookup: dict,
                       user: str, status: str, kommentar: str = "") -> int:
    now    = now_local().isoformat(timespec="seconds")
    aktion = "EINREICHUNG" if status == "ZUR_PRUEFUNG" else "ENTWURF_GESPEICHERT"

    df_fc = st.session_state["db_forecast"].copy()
    db_lookup = {}
    if not df_fc.empty:
        db_lookup = df_fc.set_index(["produktgruppe", "monat"])["menge"].to_dict()

    df_long = df_long.copy()
    df_long["Menge"] = df_long["Menge"].astype(int)

    def get_vorher(row):
        key = (row[COL_PG], row["Monat"])
        if key in db_lookup:
            return db_lookup[key]
        return original_lookup.get(key, None)

    df_long["menge_alt"] = df_long.apply(get_vorher, axis=1)
    df_long["differenz"] = df_long.apply(
        lambda r: r["Menge"] - r["menge_alt"] if r["menge_alt"] is not None else 0, axis=1
    )
    geaendert_mask = (df_long["menge_alt"].notna()) & (df_long["menge_alt"] != df_long["Menge"])
    geaendert      = int(geaendert_mask.sum())

    history_rows = []
    for _, r in df_long.iterrows():
        pg    = r[COL_PG]
        monat = r["Monat"]
        menge_neu = int(r["Menge"])
        mask = (df_fc["produktgruppe"] == pg) & (df_fc["monat"] == monat)
        if mask.any():
            df_fc.loc[mask, ["menge", "status", "ersteller", "kommentar", "updated_at"]] = [
                menge_neu, status, user, kommentar, now
            ]
        else:
            df_fc = pd.concat([df_fc, pd.DataFrame([{
                "produktgruppe": pg, "monat": monat, "menge": menge_neu,
                "status": status, "ersteller": user, "kommentar": kommentar,
                "created_at": now, "updated_at": now
            }])], ignore_index=True)

    if geaendert > 0:
        for _, r in df_long[geaendert_mask].iterrows():
            history_rows.append({
                "id":            len(st.session_state["db_history"]) + len(history_rows) + 1,
                "produktgruppe": r[COL_PG],
                "monat":         r["Monat"],
                "menge_alt":     int(r["menge_alt"]),
                "menge_neu":     int(r["Menge"]),
                "differenz":     int(r["differenz"]),
                "status":        status,
                "user":          user,
                "kommentar":     kommentar,
                "aktion":        aktion,
                "timestamp":     now,
            })

    st.session_state["db_forecast"] = df_fc
    if history_rows:
        st.session_state["db_history"] = pd.concat(
            [st.session_state["db_history"], pd.DataFrame(history_rows)],
            ignore_index=True
        )
    return geaendert


def update_status(status_neu: str, user: str, kommentar: str = "") -> int:
    now   = now_local().isoformat(timespec="seconds")
    df_fc = st.session_state["db_forecast"].copy()
    mask  = df_fc["status"] == "ZUR_PRUEFUNG"
    affected = int(mask.sum())
    df_fc.loc[mask, ["status", "updated_at"]] = [status_neu, now]
    st.session_state["db_forecast"] = df_fc

    aktion = "FREIGABE" if status_neu == "FREIGEGEBEN" else "ABLEHNUNG"
    st.session_state["db_history"] = pd.concat(
        [st.session_state["db_history"], pd.DataFrame([{
            "id":            len(st.session_state["db_history"]) + 1,
            "produktgruppe": None, "monat": None,
            "menge_alt":     None, "menge_neu": None, "differenz": None,
            "status":        status_neu, "user": user,
            "kommentar":     kommentar, "aktion": aktion, "timestamp": now,
        }])], ignore_index=True
    )
    return affected


def load_forecast() -> pd.DataFrame:
    return st.session_state["db_forecast"].copy()


def load_change_history(limit: int = 200) -> pd.DataFrame:
    df_h = st.session_state["db_history"].copy()
    if df_h.empty:
        return df_h
    df_h = df_h[
        df_h["menge_alt"].notna() &
        df_h["menge_neu"].notna() &
        (df_h["menge_alt"] != df_h["menge_neu"])
    ].sort_values("id", ascending=False).head(limit)
    return df_h


def get_current_status() -> str:
    df_db = load_forecast()
    if df_db.empty:
        return "KEINE_DATEN"
    return df_db["status"].mode().iloc[0]


def overlay_db_on_pivot(pivot: pd.DataFrame, db_df: pd.DataFrame,
                        editor_monat_strs: list, status_filter: str = None) -> pd.DataFrame:
    if db_df.empty:
        return pivot
    db = db_df.copy()
    if status_filter is not None:
        db = db[db["status"] == status_filter]
    if db.empty:
        return pivot
    db = db[db["monat"].isin(editor_monat_strs)]
    if db.empty:
        return pivot
    db = db.rename(columns=DB_TO_COL)
    pivot_long = pivot.melt(id_vars=KEY_COLS, var_name="monat", value_name="menge_orig")
    merged = pivot_long.merge(
        db[KEY_COLS + ["monat", "menge"]],
        on=KEY_COLS + ["monat"], how="left", suffixes=("", "_db"),
    )
    merged["menge_final"] = merged["menge"].fillna(merged["menge_orig"]).astype(int)
    result = merged.pivot_table(
        index=KEY_COLS, columns="monat", values="menge_final",
        aggfunc="first", fill_value=0,
    ).astype(int)
    result = result.reindex(columns=editor_monat_strs, fill_value=0)
    return result.reset_index()


def build_styled_excel(pivot_df: pd.DataFrame,
                       changed_cells: set = None,
                       history_df: pd.DataFrame = None) -> bytes:
    if changed_cells is None:
        changed_cells = set()

    buffer = BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        pivot_df.to_excel(writer, index=False, sheet_name="Prognose")
        ws = writer.sheets["Prognose"]

        n_rows, n_cols = ws.max_row, ws.max_column
        header_fill  = PatternFill("solid", start_color="000055")
        header_font  = Font(name="Arial", bold=True, color="FFFFFF", size=11)
        body_font    = Font(name="Arial", color="000055", size=10)
        changed_fill = PatternFill("solid", start_color="FFE0E0")
        changed_font = Font(name="Arial", color="000055", size=10, bold=True)
        thin   = Side(border_style="thin", color="BFBFBF")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        center = Alignment(horizontal="center", vertical="center")
        right  = Alignment(horizontal="right",  vertical="center")
        left   = Alignment(horizontal="left",   vertical="center")

        for col_idx in range(1, n_cols + 1):
            c = ws.cell(row=1, column=col_idx)
            c.fill, c.font, c.alignment, c.border = header_fill, header_font, center, border

        col_headers = [ws.cell(row=1, column=ci).value for ci in range(1, n_cols + 1)]
        n_key_cols  = len(KEY_COLS)

        for row_idx in range(2, n_rows + 1):
            pg_value = ws.cell(row=row_idx, column=1).value
            for col_idx in range(1, n_cols + 1):
                c = ws.cell(row=row_idx, column=col_idx)
                c.font, c.border = body_font, border
                if col_idx <= n_key_cols:
                    c.alignment = left
                else:
                    c.alignment    = right
                    c.number_format = "#,##0"
                    monat_header   = col_headers[col_idx - 1]
                    if (pg_value, monat_header) in changed_cells:
                        c.fill, c.font = changed_fill, changed_font

        ws.column_dimensions[get_column_letter(1)].width = 32
        for col_idx in range(n_key_cols + 1, n_cols + 1):
            ws.column_dimensions[get_column_letter(col_idx)].width = 11
        ws.freeze_panes = f"{get_column_letter(n_key_cols + 1)}2"
        ws.auto_filter.ref = ws.dimensions

        # Sheet 2: Legende
        ws2 = writer.book.create_sheet("Legende")
        download_ts = now_local().strftime("%d.%m.%Y %H:%M:%S")
        ws2["A1"] = "S&OP Prognose – Export-Information"
        ws2["A1"].font = Font(bold=True, color="000055", size=14)
        ws2["A3"] = "Erstellt am:"
        ws2["A3"].font = Font(color="000055", size=10, bold=True)
        ws2["B3"] = download_ts
        ws2["B3"].font = Font(color="000055", size=10)
        ws2["A4"] = "Zeitzone:"
        ws2["A4"].font = Font(color="000055", size=10, bold=True)
        ws2["B4"] = "Europe/Berlin"
        ws2["B4"].font = Font(color="000055", size=10)
        ws2["A5"] = "Anzahl bearbeitete Werte:"
        ws2["A5"].font = Font(color="000055", size=10, bold=True)
        ws2["B5"] = len(changed_cells)
        ws2["B5"].font = Font(color="000055", size=10)
        ws2["A7"] = "Legende Farbcodierung"
        ws2["A7"].font = Font(bold=True, color="000055", size=12)
        ws2["A9"] = ""
        ws2["A9"].fill = changed_fill
        ws2["B9"] = "Manuell durch Planer angepasster Wert"
        ws2["B9"].font = Font(color="000055", size=10, bold=True)
        ws2["B10"] = "(weicht vom ursprünglichen Forecast ab)"
        ws2["B10"].font = Font(color="666666", size=9, italic=True)
        ws2["A12"] = "Inhalt der Sheets"
        ws2["A12"].font = Font(bold=True, color="000055", size=12)
        ws2["A14"] = "Prognose"
        ws2["A14"].font = Font(color="000055", size=10, bold=True)
        ws2["B14"] = "Pivotierte Forecast-Tabelle (Produktgruppe × Monat)"
        ws2["B14"].font = Font(color="000055", size=10)
        if history_df is not None and not history_df.empty:
            ws2["A15"] = "Änderungsprotokoll"
            ws2["A15"].font = Font(color="000055", size=10, bold=True)
            ws2["B15"] = "Vollständige Historie aller Wertänderungen mit Vorher/Nachher/Bearbeiter"
            ws2["B15"].font = Font(color="000055", size=10)
        ws2.column_dimensions["A"].width = 30
        ws2.column_dimensions["B"].width = 60

        # Sheet 3: Änderungsprotokoll
        if history_df is not None and not history_df.empty:
            ws3 = writer.book.create_sheet("Änderungsprotokoll")
            history_export = history_df.copy()
            history_export["Zeitpunkt"]    = pd.to_datetime(history_export["timestamp"]).dt.strftime("%d.%m.%Y %H:%M")
            history_export["Bearbeiter"]   = history_export["user"].fillna("")
            history_export["Aktion"]       = history_export["aktion"]
            history_export["Produktgruppe"] = history_export["produktgruppe"]
            history_export["Monat"]        = history_export["monat"]
            history_export["Vorher"]       = history_export["menge_alt"]
            history_export["Nachher"]      = history_export["menge_neu"]
            history_export["Δ"]            = history_export["differenz"]
            history_export["Status"]       = history_export["status"].map(STATUS_LABEL).fillna(history_export["status"])
            history_export["Kommentar"]    = history_export["kommentar"].fillna("")
            export_cols = ["Zeitpunkt", "Bearbeiter", "Aktion", "Produktgruppe",
                           "Monat", "Vorher", "Nachher", "Δ", "Status", "Kommentar"]
            history_export = history_export[export_cols]

            for col_idx, col_name in enumerate(export_cols, start=1):
                c = ws3.cell(row=1, column=col_idx, value=col_name)
                c.fill, c.font, c.alignment, c.border = header_fill, header_font, center, border

            for row_idx, (_, hrow) in enumerate(history_export.iterrows(), start=2):
                for col_idx, col_name in enumerate(export_cols, start=1):
                    val = hrow[col_name]
                    if pd.isna(val):
                        val = ""
                    c = ws3.cell(row=row_idx, column=col_idx, value=val)
                    c.font, c.border = body_font, border
                    if col_name in ("Vorher", "Nachher", "Δ"):
                        c.alignment     = right
                        c.number_format = "#,##0"
                        if col_name == "Δ" and isinstance(val, (int, float)) and val != 0:
                            c.fill, c.font = changed_fill, changed_font
                    else:
                        c.alignment = left

            widths = {"Zeitpunkt": 18, "Bearbeiter": 20, "Aktion": 22,
                      "Produktgruppe": 32, "Monat": 10,
                      "Vorher": 12, "Nachher": 12, "Δ": 10,
                      "Status": 14, "Kommentar": 35}
            for col_idx, col_name in enumerate(export_cols, start=1):
                ws3.column_dimensions[get_column_letter(col_idx)].width = widths.get(col_name, 14)
            ws3.freeze_panes = "A2"
            ws3.auto_filter.ref = ws3.dimensions

    return buffer.getvalue()


# ============================================================
# CSS
# ============================================================
st.markdown(
    f"""
    <style>
        html, body, [class*="css"], [data-testid="stAppViewContainer"],
        [data-testid="stMarkdownContainer"], [data-testid="stMetricLabel"],
        [data-testid="stMetricValue"], [data-testid="stMetricDelta"],
        [data-testid="stSidebar"], [data-testid="stSidebar"] *,
        .stDataFrame, .stDataEditor, h1, h2, h3, h4, h5, h6,
        p, span, label, div, li, th, td, a {{
            color: {PRIMARY} !important;
        }}
        [data-baseweb="tag"] {{
            background-color: {PRIMARY} !important;
            color: #FFFFFF !important;
        }}
        [data-baseweb="tag"] * {{ color: #FFFFFF !important; }}
        .stButton > button, .stDownloadButton > button {{
            color: {PRIMARY} !important;
            background-color: #FFFFFF !important;
            border: 2px solid {PRIMARY} !important;
            font-weight: 600;
        }}
        .stButton > button *, .stDownloadButton > button * {{ color: {PRIMARY} !important; }}
        .stButton > button:hover, .stDownloadButton > button:hover {{
            background-color: {PRIMARY} !important;
            color: #FFFFFF !important;
        }}
        .stButton > button:hover *, .stDownloadButton > button:hover * {{ color: #FFFFFF !important; }}
        .stDataFrame thead tr th, .stDataEditor thead tr th {{
            background-color: {PRIMARY} !important;
            color: #FFFFFF !important;
            font-weight: 600;
        }}
        .stDataFrame thead tr th *, .stDataEditor thead tr th * {{ color: #FFFFFF !important; }}
        [data-testid="stHeader"] {{ background: transparent; height: 0; }}
        .block-container {{ padding-top: 0.5rem !important; }}
        .wp-sticky-header {{
            position: sticky; top: 0; z-index: 999;
            background-color: #FFFFFF;
            padding: 8px 0 0 0; margin: 0 0 12px 0;
            border-bottom: 6px solid {PRIMARY};
            box-shadow: 0 2px 6px rgba(0,0,0,0.08);
        }}
        .wp-sticky-header .wp-header-row {{
            display: flex; align-items: center; justify-content: space-between;
            gap: 16px; padding: 0 8px;
        }}
        .wp-sticky-header .wp-logo-left, .wp-sticky-header .wp-logo-right {{
            display: flex; align-items: center; flex: 0 0 auto;
        }}
        .wp-sticky-header .wp-logo-left img, .wp-sticky-header .wp-logo-right img {{
            max-height: {LOGO_MAXH}px; width: auto;
            image-rendering: -webkit-optimize-contrast;
            image-rendering: crisp-edges; display: block;
        }}
        .wp-sticky-header .wp-title {{ flex: 1 1 auto; text-align: center; }}
        .wp-sticky-header .wp-title h1 {{
            color: {PRIMARY} !important;
            margin: 4px 0; font-family: Arial, sans-serif;
            font-weight: 700; font-size: 2.4rem; line-height: 1.1;
        }}
    </style>
    """,
    unsafe_allow_html=True,
)


def encode_logo(path: Path) -> str:
    if not path.exists():
        return ""
    return f"data:image/png;base64,{base64.b64encode(path.read_bytes()).decode()}"


# ============================================================
# Sticky Header
# ============================================================
logo_wp_src     = encode_logo(LOGO_FILE)
logo_lemken_src = encode_logo(LEMKEN_FILE)

st.markdown(
    f"""
    <div class="wp-sticky-header">
        <div class="wp-header-row">
            <div class="wp-logo-left">
                {f'<img src="{logo_wp_src}" alt="W&P" />' if logo_wp_src else ''}
            </div>
            <div class="wp-title">
                <h1>S&amp;OP Prognose</h1>
            </div>
            <div class="wp-logo-right">
                {f'<img src="{logo_lemken_src}" alt="Lemken" />' if logo_lemken_src else ''}
            </div>
        </div>
    </div>
    """,
    unsafe_allow_html=True,
)


# ============================================================
# Daten laden
# ============================================================
@st.cache_data(show_spinner="Daten werden geladen...")
def load_data_from_upload(file_bytes: bytes) -> pd.DataFrame:
    return prepare_df(read_excel_smart(BytesIO(file_bytes)))


@st.cache_data(show_spinner=False)
def build_original_lookup(df: pd.DataFrame) -> dict:
    tmp = df.copy()
    tmp["MonatStr"] = tmp[COL_DATUM].dt.strftime("%Y-%m")
    return tmp.groupby([COL_PG, "MonatStr"], sort=False)[COL_MENGE].sum().to_dict()


if "uploaded_bytes" not in st.session_state:
    st.session_state["uploaded_bytes"] = None
if "cache_buster" not in st.session_state:
    st.session_state["cache_buster"] = 0


# ============================================================
# Sidebar
# ============================================================
st.sidebar.header("Datenquelle")

uploaded = st.sidebar.file_uploader("Excel hochladen", type=["xlsx"])
if uploaded is not None:
    new_bytes = uploaded.getvalue()
    if new_bytes != st.session_state["uploaded_bytes"]:
        st.session_state["uploaded_bytes"] = new_bytes
        st.session_state["cache_buster"]  += 1
        st.session_state["db_forecast"]    = pd.DataFrame(columns=[
            "produktgruppe", "monat", "menge", "status",
            "ersteller", "kommentar", "created_at", "updated_at"])
        st.session_state["db_history"]     = pd.DataFrame(columns=[
            "id", "produktgruppe", "monat", "menge_alt", "menge_neu",
            "differenz", "status", "user", "kommentar", "aktion", "timestamp"])
        build_original_lookup.clear()
        for k in [k for k in st.session_state if k.startswith("forecast_edit")]:
            del st.session_state[k]
        st.sidebar.success("Excel geladen.")
        st.rerun()

if not st.session_state["uploaded_bytes"]:
    st.info("👆 Bitte Excel-Datei in der Sidebar hochladen.")
    st.stop()

try:
    df = load_data_from_upload(st.session_state["uploaded_bytes"])
except Exception as e:
    st.error(f"Hochgeladene Datei konnte nicht gelesen werden:\n\n{e}")
    st.stop()

ORIGINAL_LOOKUP = build_original_lookup(df)

heute              = pd.Timestamp(now_local().date())
heute_monatsanfang = pd.Timestamp(year=heute.year, month=heute.month, day=1)
cutoff_monat       = heute_monatsanfang - pd.DateOffset(months=LOOKBACK_MONTHS - 1)
EDITOR_DATES       = sorted([d for d in df[COL_DATUM].unique() if pd.Timestamp(d) >= cutoff_monat])
EDITOR_MONAT_STRS  = [pd.Timestamp(d).strftime("%Y-%m") for d in EDITOR_DATES]
DATA_START_DATE    = df[COL_DATUM].min()

st.sidebar.markdown("---")
st.sidebar.header("Benutzer")
rolle = st.sidebar.radio("Rolle", ["Planer", "Reviewer"], index=0)
user  = st.sidebar.text_input("Name", value="", placeholder="Vor- und Nachname")

st.sidebar.markdown("---")
st.sidebar.header("Status")
current_status = get_current_status()
st.sidebar.markdown(
    f"""
    <div style="padding:8px 12px; background-color:{STATUS_COLOR.get(current_status, GRAY_MEDIUM)};
                color:white; border-radius:4px; font-weight:600; text-align:center;">
        {STATUS_LABEL.get(current_status, "Keine Daten")}
    </div>
    """,
    unsafe_allow_html=True,
)

st.sidebar.markdown("---")
st.sidebar.header("Filter")

pg_opts  = sorted(df[COL_PG].unique())
sel_pg   = st.sidebar.multiselect("Produktgruppe", pg_opts, default=pg_opts)
jahre    = sorted(df["Jahr"].unique())
sel_jahr = st.sidebar.multiselect("Jahr", jahre, default=jahre)
monate   = list(range(1, 13))
sel_monat = st.sidebar.select_slider("Monatsbereich", options=monate, value=(min(monate), max(monate)))

phase_opts_all = list(df[COL_PHASE].unique())
phase_opts     = [p for p in phase_opts_all if p in (PHASE_IST, PHASE_FORECAST)]
phase_opts     = sorted(phase_opts) if phase_opts else sorted(phase_opts_all)
sel_phase      = st.sidebar.multiselect("Phase", phase_opts, default=phase_opts)


# ============================================================
# Filter anwenden
# ============================================================
mask = (
    df[COL_PG].isin(sel_pg)
    & df["Jahr"].isin(sel_jahr)
    & df["MonatNr"].between(sel_monat[0], sel_monat[1])
    & df[COL_PHASE].isin(sel_phase)
)
df_f = df.loc[mask].copy()

if df_f.empty:
    st.warning("Keine Daten für die aktuelle Filterauswahl.")
    st.stop()


# ============================================================
# KPIs + Donut
# ============================================================
mask_ist      = df_f[COL_PHASE] == PHASE_IST
mask_forecast = df_f[COL_PHASE] == PHASE_FORECAST

absatz_ist      = df_f.loc[mask_ist, COL_MENGE].sum()
absatz_forecast = df_f.loc[mask_forecast, COL_MENGE].sum()
mean_pro_monat  = df_f.groupby(COL_DATUM)[COL_MENGE].sum().mean()
n_pg            = df_f[COL_PG].nunique()

last3_start   = heute_monatsanfang - pd.DateOffset(months=3)
mask_last3    = (df_f[COL_DATUM] >= last3_start) & (df_f[COL_DATUM] < heute_monatsanfang) & mask_ist
last3_monthly = df_f.loc[mask_last3].groupby(COL_DATUM)[COL_MENGE].sum()
mean_last3    = last3_monthly.mean() if not last3_monthly.empty else 0

next3_end     = heute_monatsanfang + pd.DateOffset(months=3)
mask_next3    = (df_f[COL_DATUM] >= heute_monatsanfang) & (df_f[COL_DATUM] < next3_end) & mask_forecast
next3_monthly = df_f.loc[mask_next3].groupby(COL_DATUM)[COL_MENGE].sum()
mean_next3    = next3_monthly.mean() if not next3_monthly.empty else 0

delta_runrate_str = None
if mean_last3 > 0:
    delta_runrate_pct = (mean_next3 - mean_last3) / mean_last3 * 100
    delta_runrate_str = f"{delta_runrate_pct:+.1f} % vs. Ist"


def _yoy_delta_pct(current_sum, df_full, mask_current, periode_label):
    if current_sum <= 0:
        return None
    df_current = df_full.loc[mask_current]
    if df_current.empty:
        return None
    min_d, max_d = df_current[COL_DATUM].min(), df_current[COL_DATUM].max()
    prev = df_full[
        (df_full[COL_DATUM] >= min_d - pd.DateOffset(years=1)) &
        (df_full[COL_DATUM] <= max_d - pd.DateOffset(years=1)) &
        (df_full[COL_PG].isin(sel_pg))
    ]
    prev_sum = prev[COL_MENGE].sum()
    if prev_sum <= 0:
        return None
    return f"{(current_sum - prev_sum) / prev_sum * 100:+.1f} % YoY"


yoy_ist  = _yoy_delta_pct(absatz_ist,      df, mask & mask_ist,      "Ist")
yoy_fc   = _yoy_delta_pct(absatz_forecast, df, mask & mask_forecast, "Forecast")
yoy_mean = _yoy_delta_pct(
    mean_pro_monat * df_f[COL_DATUM].nunique() if df_f[COL_DATUM].nunique() else 0,
    df, mask, "Ø/Monat",
)

start_label  = pd.Timestamp(df_f[COL_DATUM].min()).strftime("%b %Y")
col_kpi, col_pie = st.columns([3, 2], gap="large")

with col_kpi:
    k1, k2, k3 = st.columns(3)
    k4, k5, k6 = st.columns(3)
    k1.metric(f"Ist-Absatz seit {start_label} [Stk.]", fmt_int(absatz_ist),      delta=yoy_ist)
    k2.metric("Absatzprognose [Stk.]",                  fmt_int(absatz_forecast), delta=yoy_fc)
    k3.metric("Ø pro Monat [Stk.]",                     fmt_int(mean_pro_monat),  delta=yoy_mean)
    k4.metric("Produktgruppen [#]",                     fmt_int(n_pg))
    k5.metric("Ø letzte 3 Monate [Stk.]",               fmt_int(mean_last3))
    k6.metric("Ø nächste 3 Monate [Stk.]",              fmt_int(mean_next3),      delta=delta_runrate_str)

with col_pie:
    pg_agg = (df_f.groupby(COL_PG, as_index=False, sort=False)[COL_MENGE]
              .sum().sort_values(COL_MENGE, ascending=False))
    if len(pg_agg) > 10:
        top10 = pg_agg.head(10).copy()
        rest  = pg_agg.iloc[10:][COL_MENGE].sum()
        pg_display = pd.concat([top10, pd.DataFrame([{COL_PG: "Sonstige", COL_MENGE: rest}])],
                                ignore_index=True)
    else:
        pg_display = pg_agg
    fig_pie = px.pie(pg_display, names=COL_PG, values=COL_MENGE,
                     hole=0.55, color_discrete_sequence=WP_PALETTE)
    fig_pie.update_traces(textposition="inside", textinfo="percent",
                          textfont=dict(color="#FFFFFF", size=10),
                          marker=dict(line=dict(color="#FFFFFF", width=2)))
    fig_pie.update_layout(height=180, margin=dict(l=0, r=0, t=0, b=0),
                          legend=dict(orientation="v", yanchor="middle", y=0.5,
                                      xanchor="left", x=1.02,
                                      font=dict(color=PRIMARY, family="Arial", size=9)),
                          font=dict(color=PRIMARY, family="Arial"))
    st.plotly_chart(fig_pie, use_container_width=True)

st.markdown("---")


# ============================================================
# Liniendiagramm – Ist/Forecast mit Range-Slider
# ============================================================
st.subheader("Absatzverlauf – Ist & Forecast")

n_pg_selected = len(sel_pg)
DETAIL_LIMIT  = 8

col_view, col_info = st.columns([1, 3])
with col_view:
    view_mode = st.radio("Ansicht", ["Aggregat", "Pro Produktgruppe"],
                         horizontal=True, label_visibility="collapsed", key="lineplot_view_mode")
with col_info:
    if view_mode == "Pro Produktgruppe" and n_pg_selected > DETAIL_LIMIT:
        st.caption(f"⚠️ {n_pg_selected} Produktgruppen ausgewählt – zeigt Top {DETAIL_LIMIT} nach Volumen.")

fig = go.Figure()

DETAIL_PALETTE = [PRIMARY, BLUE_MEDIUM, "#4A6FA5", "#7A9BC9", "#3D5A8A",
                  "#5C7DB0", "#85A5CC", GRAY_MEDIUM]


def _add_phase_traces_per_pg(fig, df_phase, phase_name, pg_color_map, dash_style, opacity=1.0):
    for pg in pg_color_map.keys():
        sub = df_phase[df_phase[COL_PG] == pg].sort_values(COL_DATUM)
        if sub.empty:
            continue
        fig.add_trace(go.Scatter(
            x=sub[COL_DATUM], y=sub[COL_MENGE],
            mode="lines+markers", name=f"{pg} – {phase_name}",
            legendgroup=pg,
            line=dict(color=pg_color_map[pg], width=2, dash=dash_style),
            marker=dict(color=pg_color_map[pg], size=5),
            opacity=opacity,
            hovertemplate=(f"<b>{pg}</b><br>{phase_name}<br>"
                           "%{x|%b %Y}<br>%{y:,.0f} Stk.<extra></extra>"),
        ))


if view_mode == "Aggregat":
    agg_total     = (df_f.groupby([COL_DATUM, COL_PHASE], as_index=False, sort=True)[COL_MENGE]
                     .sum().sort_values(COL_DATUM))
    ist_data      = agg_total[agg_total[COL_PHASE] == PHASE_IST].sort_values(COL_DATUM)
    forecast_data = agg_total[agg_total[COL_PHASE] == PHASE_FORECAST].sort_values(COL_DATUM)

    if not ist_data.empty and not forecast_data.empty:
        last_ist = ist_data.iloc[[-1]].copy()
        last_ist[COL_PHASE] = PHASE_FORECAST
        forecast_data = pd.concat([last_ist, forecast_data], ignore_index=True)

    if not ist_data.empty:
        fig.add_trace(go.Scatter(x=ist_data[COL_DATUM], y=ist_data[COL_MENGE],
                                  mode="lines+markers", name="Ist",
                                  line=dict(color=PRIMARY, width=2.5),
                                  marker=dict(color=PRIMARY, size=6),
                                  hovertemplate="<b>Ist</b><br>%{x|%b %Y}<br>%{y:,.0f} Stk.<extra></extra>"))
    if not forecast_data.empty:
        fig.add_trace(go.Scatter(x=forecast_data[COL_DATUM], y=forecast_data[COL_MENGE],
                                  mode="lines+markers", name="Forecast",
                                  line=dict(color=FORECAST_GRAY, width=2.5),
                                  marker=dict(color=FORECAST_GRAY, size=6),
                                  hovertemplate="<b>Forecast</b><br>%{x|%b %Y}<br>%{y:,.0f} Stk.<extra></extra>"))
    agg_total_for_range = agg_total
else:
    pg_volumes   = df_f.groupby(COL_PG)[COL_MENGE].sum().sort_values(ascending=False)
    top_pgs      = pg_volumes.head(DETAIL_LIMIT).index.tolist()
    pg_color_map = {pg: DETAIL_PALETTE[i % len(DETAIL_PALETTE)] for i, pg in enumerate(top_pgs)}
    df_detail    = df_f[df_f[COL_PG].isin(top_pgs)]
    df_ist       = df_detail[df_detail[COL_PHASE] == PHASE_IST]
    df_forecast  = df_detail[df_detail[COL_PHASE] == PHASE_FORECAST]

    if not df_ist.empty and not df_forecast.empty:
        bridges = []
        for pg in top_pgs:
            sub_ist = df_ist[df_ist[COL_PG] == pg].sort_values(COL_DATUM)
            if sub_ist.empty:
                continue
            bridge = sub_ist.iloc[[-1]].copy()
            bridge[COL_PHASE] = PHASE_FORECAST
            bridges.append(bridge)
        if bridges:
            df_forecast = pd.concat([pd.concat(bridges), df_forecast], ignore_index=True)

    _add_phase_traces_per_pg(fig, df_ist,      "Ist",      pg_color_map, dash_style="solid")
    _add_phase_traces_per_pg(fig, df_forecast, "Forecast", pg_color_map, dash_style="dash", opacity=0.85)
    agg_total_for_range = df_detail.groupby(COL_DATUM, as_index=False)[COL_MENGE].sum()

if not agg_total_for_range.empty:
    x_min, x_max = agg_total_for_range[COL_DATUM].min(), agg_total_for_range[COL_DATUM].max()
else:
    x_min, x_max = heute - pd.DateOffset(years=1), heute + pd.DateOffset(years=1)

x_axis_min    = min(x_min, heute) - pd.Timedelta(days=15)
x_axis_max    = max(x_max, heute) + pd.Timedelta(days=15)
default_start = max(x_axis_min, heute - pd.DateOffset(years=3))

fig.update_layout(
    height=520, margin=dict(l=10, r=10, t=20, b=10),
    legend=dict(orientation="h", yanchor="bottom", y=-0.30, xanchor="left", x=0,
                font=dict(color=PRIMARY, family="Arial", size=11)),
    xaxis_title=None, yaxis_title="Menge", plot_bgcolor="white",
    font=dict(color=PRIMARY, family="Arial"),
    hovermode="x unified" if view_mode == "Aggregat" else "closest",
)
fig.update_xaxes(
    showgrid=True, gridcolor=GRAY_LIGHT,
    tickfont=dict(color=PRIMARY), title_font=dict(color=PRIMARY),
    range=[default_start, x_axis_max],
    rangeslider=dict(visible=True, thickness=0.06, bgcolor=BLUE_LIGHTEST),
    rangeselector=dict(
        buttons=[
            dict(count=1, label="1J", step="year", stepmode="backward"),
            dict(count=3, label="3J", step="year", stepmode="backward"),
            dict(count=5, label="5J", step="year", stepmode="backward"),
            dict(step="all", label="Alle"),
        ],
        bgcolor="#FFFFFF", bordercolor=PRIMARY, borderwidth=1,
        font=dict(color=PRIMARY, size=10), x=0, y=1.10,
    ),
)
fig.update_yaxes(showgrid=True, gridcolor=GRAY_LIGHT,
                 tickfont=dict(color=PRIMARY), title_font=dict(color=PRIMARY), tickformat=",")

fig.add_shape(type="line", x0=heute, x1=heute, y0=0, y1=1,
              yref="paper", xref="x",
              line=dict(color=ACCENT_RED, width=2, dash="dash"), layer="above")
fig.add_annotation(x=heute, y=1.02, xref="x", yref="paper",
                   text=heute.strftime("%d.%m.%y"), showarrow=False,
                   font=dict(color=ACCENT_RED, size=11, family="Arial"),
                   bgcolor="rgba(255,255,255,0.9)",
                   bordercolor=ACCENT_RED, borderwidth=1, borderpad=3)

st.plotly_chart(fig, use_container_width=True)
st.markdown("---")


# ============================================================
# Prognose-Editor
# ============================================================
st.subheader("Prognose (editierbar)")

if EDITOR_DATES:
    von = pd.Timestamp(EDITOR_DATES[0]).strftime('%b %Y')
    bis = pd.Timestamp(EDITOR_DATES[-1]).strftime('%b %Y')
    st.caption(f"Sichtbar: {LOOKBACK_MONTHS} Monate Rückblick + Forecast ({von} – {bis})")
else:
    st.caption("Keine Daten im Editor-Zeitraum verfuegbar.")

df_editor_scope = df_f[df_f[COL_DATUM].isin(EDITOR_DATES)]

if df_editor_scope.empty:
    st.info("Keine Daten im sichtbaren Zeitraum.")
else:
    pivot = (df_editor_scope.pivot_table(
                index=KEY_COLS, columns=COL_DATUM,
                values=COL_MENGE, aggfunc="sum", fill_value=0)
             .astype(int))
    pivot.columns = [pd.Timestamp(d).strftime("%Y-%m") for d in pivot.columns]
    pivot = pivot.reindex(columns=EDITOR_MONAT_STRS, fill_value=0).reset_index()

    db_df = load_forecast()
    pivot = overlay_db_on_pivot(pivot, db_df, EDITOR_MONAT_STRS)

    state_key = "forecast_edit"
    reset_key = f"{state_key}_signature"
    signature = (tuple(sel_pg), tuple(sel_jahr), sel_monat,
                 tuple(sel_phase), tuple(EDITOR_MONAT_STRS),
                 st.session_state["cache_buster"])

    if state_key not in st.session_state or st.session_state.get(reset_key) != signature:
        st.session_state[state_key] = pivot.copy()
        st.session_state[reset_key] = signature

    heute_str        = heute_monatsanfang.strftime("%Y-%m")
    historisch_strs  = [m for m in EDITOR_MONAT_STRS if m <= heute_str]
    forecast_strs    = [m for m in EDITOR_MONAT_STRS if m >  heute_str]

    DE_MONATE = ["Jan", "Feb", "Mär", "Apr", "Mai", "Jun",
                 "Jul", "Aug", "Sep", "Okt", "Nov", "Dez"]

    def _monat_label(monat_str: str) -> str:
        y, m = monat_str.split("-")
        return f"{DE_MONATE[int(m) - 1]} {y}"

    def _row_change_summary(edited_df, original_lookup, monat_cols):
        summary = {}
        for _, row in edited_df.iterrows():
            pg = row[COL_PG]
            n_chg, delta_sum = 0, 0
            for m in monat_cols:
                neu = int(row[m]) if pd.notna(row[m]) else 0
                alt = original_lookup.get((pg, m), 0)
                try:
                    alt = int(alt)
                except (TypeError, ValueError):
                    alt = 0
                if neu != alt:
                    n_chg += 1
                    delta_sum += (neu - alt)
            summary[pg] = (n_chg, delta_sum)
        return summary

    current_state = st.session_state[state_key]
    row_summary   = _row_change_summary(current_state, ORIGINAL_LOOKUP, EDITOR_MONAT_STRS)

    def _marker_text(pg):
        n, _ = row_summary.get(pg, (0, 0))
        return f"🔴 {n}" if n > 0 else ""

    display_state = current_state.copy()
    display_state.insert(0, "Δ", display_state[COL_PG].map(_marker_text))

    disabled_reviewer = (rolle == "Reviewer")

    column_config = {
        "Δ": st.column_config.TextColumn("Δ", disabled=True, pinned="left", width="small",
                                          help="Anzahl Änderungen in dieser Zeile"),
        COL_PG: st.column_config.TextColumn(COL_PG, disabled=True, pinned="left", width="medium"),
    }
    for m in historisch_strs:
        column_config[m] = st.column_config.NumberColumn(
            f"🔒 {_monat_label(m)}", disabled=True, format="%d",
            help="Vergangenheitswert – nicht editierbar")
    for m in forecast_strs:
        column_config[m] = st.column_config.NumberColumn(
            _monat_label(m), disabled=disabled_reviewer, format="%d",
            help="Forecast – editierbar", min_value=0, step=1)

    st.markdown(
        f"""
        <style>
        [data-testid="stDataEditor"] [aria-readonly="true"][role="gridcell"] {{
            background-color: #F0F0F2 !important;
            color: #555555 !important;
        }}
        [data-testid="stDataEditor"] [role="gridcell"][aria-colindex="1"] {{
            background-color: #FFFFFF !important;
            font-weight: 700 !important;
            color: {ACCENT_RED} !important;
            text-align: center !important;
        }}
        [data-testid="stDataEditor"] [role="gridcell"][aria-colindex="2"] {{
            background-color: #FAFAFC !important;
            font-weight: 600 !important;
        }}
        </style>
        """,
        unsafe_allow_html=True,
    )

    edited_with_marker = st.data_editor(
        display_state, num_rows="fixed", use_container_width=True,
        height=480, disabled=disabled_reviewer, key="editor_main",
        column_config=column_config,
    )
    edited = edited_with_marker.drop(columns=["Δ"])
    st.session_state[state_key] = edited

    # Live-Diff
    def _compute_changes(edited_df, original_lookup, monat_cols):
        changes = []
        for _, row in edited_df.iterrows():
            pg = row[COL_PG]
            for m in monat_cols:
                neu = int(row[m]) if pd.notna(row[m]) else 0
                alt = original_lookup.get((pg, m), 0)
                try:
                    alt = int(alt)
                except (TypeError, ValueError):
                    alt = 0
                if neu != alt:
                    changes.append((pg, m, alt, neu))
        return changes

    changes = _compute_changes(edited, ORIGINAL_LOOKUP, EDITOR_MONAT_STRS)

    if changes:
        n_changes = len(changes)
        st.markdown(
            f"""
            <div style="padding:8px 12px; margin: 8px 0;
                        background-color:#FFE0E0; border-left: 4px solid {ACCENT_RED};
                        border-radius:4px; color:{PRIMARY}; font-weight:600;">
                ✏️ {n_changes} Wert{'e' if n_changes != 1 else ''} bearbeitet
                <span style="font-weight:400; font-size: 0.9em;">
                    – im Excel-Export werden diese Zellen hellrot hinterlegt
                </span>
            </div>
            """,
            unsafe_allow_html=True,
        )
        with st.expander(f"Änderungen anzeigen ({n_changes})"):
            df_changes = pd.DataFrame(changes, columns=[COL_PG, "Monat", "Vorher", "Nachher"])
            df_changes["Δ"]   = df_changes["Nachher"] - df_changes["Vorher"]
            df_changes["Δ %"] = df_changes.apply(
                lambda r: f"{(r['Δ'] / r['Vorher'] * 100):+.1f} %" if r["Vorher"] != 0 else "n.a.",
                axis=1)
            df_changes["Monat"] = df_changes["Monat"].apply(_monat_label)
            st.dataframe(df_changes, use_container_width=True, hide_index=True, height=200)
        st.session_state["forecast_changes"] = {(pg, m) for pg, m, _, _ in changes}
    else:
        st.session_state["forecast_changes"] = set()

    # Aktionen
    st.markdown("##### Aktionen")
    kommentar = st.text_input("Kommentar (optional)", key="kommentar_input")

    if rolle == "Planer":
        col_a, col_b = st.columns(2)
        with col_a:
            if st.button("💾 Als Entwurf speichern", use_container_width=True, disabled=not user):
                df_long = edited.melt(id_vars=KEY_COLS, var_name="Monat", value_name="Menge")
                n = save_forecast_bulk(df_long, ORIGINAL_LOOKUP,
                                       user=user, status="ENTWURF", kommentar=kommentar)
                st.success(f"Als Entwurf gespeichert ({n} Wertänderungen).")
                st.rerun()
        with col_b:
            if st.button("📤 Zur Prüfung einreichen", use_container_width=True, disabled=not user):
                df_long = edited.melt(id_vars=KEY_COLS, var_name="Monat", value_name="Menge")
                n = save_forecast_bulk(df_long, ORIGINAL_LOOKUP,
                                       user=user, status="ZUR_PRUEFUNG", kommentar=kommentar)
                st.success(f"Zur Prüfung eingereicht ({n} Wertänderungen).")
                st.rerun()
    else:
        if current_status != "ZUR_PRUEFUNG":
            st.info("Aktuell liegt keine Einreichung zur Prüfung vor.")
        else:
            col_a, col_b = st.columns(2)
            with col_a:
                if st.button("✅ Freigeben", use_container_width=True, disabled=not user):
                    n = update_status("FREIGEGEBEN", user=user, kommentar=kommentar)
                    st.success(f"{n} Einträge freigegeben.")
                    st.rerun()
            with col_b:
                if st.button("❌ Ablehnen", use_container_width=True, disabled=not user):
                    n = update_status("ABGELEHNT", user=user, kommentar=kommentar)
                    st.warning(f"{n} Einträge abgelehnt.")
                    st.rerun()

    # Excel-Export
    st.markdown("##### Export")
    if current_status == "FREIGEGEBEN":
        df_export_scope = df[df[COL_DATUM].isin(EDITOR_DATES)]
        pivot_full = (df_export_scope.groupby(KEY_COLS + [COL_DATUM], as_index=False, sort=False)[COL_MENGE].sum()
                      .pivot_table(index=KEY_COLS, columns=COL_DATUM,
                                   values=COL_MENGE, aggfunc="sum", fill_value=0).astype(int))
        pivot_full.columns = [pd.Timestamp(d).strftime("%Y-%m") for d in pivot_full.columns]
        pivot_full = pivot_full.reindex(columns=EDITOR_MONAT_STRS, fill_value=0).reset_index()
        db_df_full = load_forecast()
        pivot_full = overlay_db_on_pivot(pivot_full, db_df_full, EDITOR_MONAT_STRS,
                                         status_filter="FREIGEGEBEN")
        pivot_full = pivot_full.sort_values(KEY_COLS)

        changed_cells_export = set()
        if not db_df_full.empty:
            db_freigegeben = db_df_full[db_df_full["status"] == "FREIGEGEBEN"]
            for _, row in db_freigegeben.iterrows():
                pg, monat, neu = row["produktgruppe"], row["monat"], int(row["menge"])
                alt = ORIGINAL_LOOKUP.get((pg, monat), 0)
                try:
                    alt = int(alt)
                except (TypeError, ValueError):
                    alt = 0
                if neu != alt:
                    changed_cells_export.add((pg, monat))

        history_for_export = load_change_history(limit=1000)
        excel_bytes = build_styled_excel(pivot_full,
                                         changed_cells=changed_cells_export,
                                         history_df=history_for_export)

        if changed_cells_export:
            st.caption(f"📌 {len(changed_cells_export)} bearbeitete Werte werden im Excel hellrot hinterlegt.")

        ts = now_local().strftime("%d%m%Y_%H_%M")
        st.download_button(
            "⬇️ Freigegebene Prognose als Excel herunterladen",
            data=excel_bytes,
            file_name=f"w&p_lemken_s&op_{ts}.xlsx",
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )
    else:
        st.caption(f"Excel-Export ist erst nach Freigabe verfügbar. "
                   f"Aktueller Status: **{STATUS_LABEL.get(current_status, 'Keine Daten')}**")


# ============================================================
# Änderungshistorie
# ============================================================
with st.expander("Änderungshistorie"):
    hist = load_change_history()
    if hist.empty:
        st.caption("Noch keine Wertänderungen protokolliert.")
    else:
        hist_view = hist.copy()
        hist_view["Zeitpunkt"]  = pd.to_datetime(hist_view["timestamp"]).dt.strftime("%d.%m.%Y %H:%M")
        hist_view["Bearbeiter"] = hist_view["user"]
        hist_view["Aktion"]     = hist_view["aktion"]
        hist_view[COL_PG]       = hist_view["produktgruppe"]
        hist_view["Monat"]      = hist_view["monat"]
        hist_view["Vorher"]     = hist_view["menge_alt"].astype("Int64")
        hist_view["Nachher"]    = hist_view["menge_neu"].astype("Int64")
        hist_view["Δ"]          = hist_view["differenz"].astype("Int64")
        hist_view["Status"]     = hist_view["status"].map(STATUS_LABEL).fillna(hist_view["status"])
        hist_view["Kommentar"]  = hist_view["kommentar"].fillna("")
        cols = ["Zeitpunkt", "Bearbeiter", "Aktion", COL_PG,
                "Monat", "Vorher", "Nachher", "Δ", "Status", "Kommentar"]
        st.dataframe(hist_view[cols], use_container_width=True, height=350, hide_index=True)